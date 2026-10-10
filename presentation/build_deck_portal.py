"""
Vienotais Zemes Reģistrs — portāla prezentācija (latvija.gov.lv krāsas, LV).

Palaišana:
    .venv/bin/python build_deck_portal.py

Ievade:  clips/*.mp4, posters/*.png
Izvade:  Vienotais-Zemes-Registrs-portals.pptx
"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
CLIPS = HERE / "clips"
POSTERS = HERE / "posters"

# ── latvija.gov.lv / lietotnes palete (src/index.css)
BG = RGBColor(0xF4, 0xF7, 0xFB)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xE9, 0xF2, 0xFC)
INK = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x4A, 0x4E, 0x56)
FAINT = RGBColor(0x5C, 0x65, 0x70)
LINE = RGBColor(0xC1, 0xD2, 0xE5)
ACCENT = RGBColor(0xA0, 0x1B, 0x36)
ACCENT_DEEP = RGBColor(0x7D, 0x15, 0x29)
MAP_BLUE = RGBColor(0x2B, 0x6C, 0xB0)
GREEN = RGBColor(0x1B, 0x6B, 0x38)
ORANGE = RGBColor(0xC4, 0x5C, 0x00)
FONT = "Ubuntu"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = prs.slide_width
H = prs.slide_height
BLANK = prs.slide_layouts[6]


def new_slide(notes: str):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    s.notes_slide.notes_text_frame.text = notes
    trans = etree.SubElement(s._element, qn("p:transition"), spd="med")
    etree.SubElement(trans, qn("p:fade"))
    clr = s._element.find(qn("p:clrMapOvr"))
    if clr is not None:
        clr.addnext(trans)
    return s


def text(
    slide,
    x,
    y,
    w,
    h,
    content,
    size=20,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    spacing=None,
    tracking=None,
):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = para
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        if tracking:
            r._r.get_or_add_rPr().set("spc", str(tracking))
    return tb


def card(slide, x, y, w, h, fill=SURFACE, radius=0.06, line=LINE):
    """Opaque portal card — white surface, hairline border."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def accent_bar(slide, x, y, w=0.08, h=0.9):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = ACCENT
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def pill(slide, x, y, w, label, color=ACCENT, size=12):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    shp.adjustments[0] = 0.5
    shp.fill.solid()
    shp.fill.fore_color.rgb = TINT
    shp.line.color.rgb = color
    shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    return shp


def eyebrow(slide, label, x=0.8, y=0.55):
    text(slide, x, y, 10, 0.28, label.upper(), size=11, bold=True, color=ACCENT, tracking=160)


def title(slide, content, x=0.8, y=0.85, w=11.7, size=36):
    text(slide, x, y, w, 1.35, content, size=size, bold=True, color=INK)


def video(slide, name, x, y, w, h):
    """Embedded clip; poster is the mid-frame PNG from convert.sh."""
    card(slide, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=SURFACE, radius=0.04)
    mp4 = CLIPS / f"{name}.mp4"
    poster = POSTERS / f"{name}.png"
    if not mp4.exists():
        text(slide, x, y + h / 2 - 0.2, w, 0.4, f"[nav klipa: {name}]", size=14, color=MUTED, align=PP_ALIGN.CENTER)
        return None
    return slide.shapes.add_movie(
        str(mp4),
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        poster_frame_image=str(poster) if poster.exists() else None,
        mime_type="video/mp4",
    )


def bullets(slide, x, y, w, items, size=17, gap=0.55, dot=ACCENT, color=INK):
    for i, item in enumerate(items):
        yy = y + i * gap
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy + 0.1), Inches(0.11), Inches(0.11))
        d.fill.solid()
        d.fill.fore_color.rgb = dot
        d.line.fill.background()
        d.shadow.inherit = False
        text(slide, x + 0.28, yy, w - 0.28, gap, item, size=size, color=color)


def footer(slide, label="VIENOTAIS ZEMES REĢISTRS"):
    text(slide, 0.8, 7.05, 6, 0.25, label, size=9, bold=True, color=FAINT, tracking=180)
    text(slide, 9.5, 7.05, 3.0, 0.25, "latvija.gov.lv stilā", size=9, color=FAINT, align=PP_ALIGN.RIGHT)


# ───────────────────────── 1. Tituls ─────────────────────────
s = new_slide(
    "[0:00–0:25] Sveiki. Vienotais Zemes Reģistrs — karte, kurā uzklikšķini uz jebkura zemes gabala Latvijā "
    "un redzi to, ko valsts jau publicējusi atvērti. Bez pieteikšanās. Šodien — trīs skatpunkti: "
    "pirms pirkuma, īpašnieks meklē atbalstu, un valsts pārvalde atlasei."
)
accent_bar(s, 0.8, 2.35, 0.1, 1.6)
text(s, 1.1, 2.2, 11, 1.0, "Vienotais Zemes Reģistrs", size=48, bold=True)
text(s, 1.1, 3.35, 10, 0.55, "Uzklikšķini zemei. Redzi visu.", size=26, color=ACCENT_DEEP)
text(
    s,
    1.1,
    4.3,
    10,
    0.8,
    ["19 atvērtie avoti · 6 panela sekcijas", "Pircējs · īpašnieks · valsts pārvalde"],
    size=16,
    color=MUTED,
    spacing=6,
)
text(s, 1.1, 6.2, 10, 0.35, "ATVĒRTO DATU HAKATONS  ·  LATVIJA", size=12, bold=True, color=FAINT, tracking=160)
footer(s)

# ───────────────────────── 2. Problēma ─────────────────────────
s = new_slide(
    "[0:25–1:00] Dati ir atvērti — bet sadalīti. Pircējs pārbauda plūdus vienā vietnē, dabu citā. "
    "Īpašnieks meklē CFLA un EKII atsevišķi. Ierēdnis atlasei saliek īpašnieka veidu, ES projektus un "
    "atvērtās atlases no dažādiem portāliem. Vienas ainas nav."
)
eyebrow(s, "Problēma")
title(s, ["Dati ir atvērti.", "Vienas ainas nav."], size=40)
for i, (big, small, who) in enumerate(
    [
        ("15+", "atsevišķas vietnes", "kādreiz jāapmeklē"),
        ("3", "lietotāju tipiem", "pircējs · īpašnieks · pārvalde"),
        ("1", "klikšķis tagad", "viss panelī"),
    ]
):
    x = 0.8 + i * 4.05
    card(s, x, 3.55, 3.8, 2.85, fill=SURFACE)
    text(s, x + 0.3, 3.75, 3.2, 0.9, big, size=52, bold=True, color=ACCENT if i < 2 else MAP_BLUE)
    text(s, x + 0.3, 4.75, 3.2, 0.45, small, size=16, bold=True)
    text(s, x + 0.3, 5.35, 3.2, 0.7, who, size=14, color=MUTED)
footer(s)

# ───────────────────────── 3. Risinājums + klikšķis ─────────────────────────
s = new_slide(
    "[1:00–2:20] Produkts. Uzklikšķinu zemei — kontūra un kadastra numurs. Rindas piepildās pa vienai, "
    "tiklīdz avots atbild. Zem katras — avots, laiks, Avots un Pieprasījums. Īpašums, ierobežojumi, "
    "plānošana, būvniecība, ES finansējums, uzņēmumi."
)
eyebrow(s, "Risinājums")
text(s, 0.8, 0.9, 4.3, 1.4, ["Viens", "klikšķis."], size=40, bold=True)
bullets(
    s,
    0.8,
    2.6,
    4.0,
    [
        "Kontūra ~1 sekundē.",
        "Rindas piepildās pa vienai.",
        "Katrai — avots un pieprasījums.",
        "Bez pieteikšanās un atslēgām.",
    ],
    size=15,
    gap=0.7,
)
video(s, "klikskis", 5.15, 0.95, 7.6, 4.75)
text(s, 5.15, 5.9, 7.6, 0.35, "Kaļķu iela 1, Rīga · 01000070142 · satelīts · CFLA Atbilst", size=11, color=MUTED)
footer(s)

# ───────────────────────── 4. Cik vietnes ─────────────────────────
s = new_slide(
    "[2:20–3:00] Cik vietnes tas apvieno. Sources matricā — 19 atvērtie avoti: VZD, TAPIS, LVĢMC, Ozols, "
    "BIS, CFLA, EKII/LVIF, LIAA, UR, VID, FM, EIS un citi. Panelī sešas sekcijas. Viena karte."
)
eyebrow(s, "Vienota vieta")
title(s, "19 avoti. 6 sekcijas. 1 karte.", size=34)
sections = [
    ("Īpašums", "VZD kadastrs, adrese, tiesības"),
    ("Ierobežojumi", "ATIS, plūdi, Ozols, NKMP"),
    ("Plānošana", "TAPIS zonējums un dokuments"),
    ("Būvniecība", "BIS lietas un būves"),
    ("ES finansējums", "ES projekti, CFLA, EKII, LIAA"),
    ("Uzņēmumi", "UR, VID, de minimis, EIS"),
]
for i, (a, b) in enumerate(sections):
    x = 0.8 + (i % 3) * 4.05
    y = 2.45 + (i // 3) * 2.15
    card(s, x, y, 3.85, 1.95)
    accent_bar(s, x + 0.2, y + 0.35, 0.07, 0.55)
    text(s, x + 0.45, y + 0.3, 3.1, 0.45, a, size=18, bold=True)
    text(s, x + 0.45, y + 0.9, 3.1, 0.8, b, size=13, color=MUTED)
footer(s)

# ───────────────────────── 5. Skatpunkts A: pircējs ─────────────────────────
s = new_slide(
    "[3:00–4:20] Skatpunkts A — pirms pirkuma. Vai gabals applūst? Vai ir ĪADT? Kas atļauts zonējumā? "
    "Slāņi un rindas nāk no tiem pašiem valsts servisiem. Vecāķu piemērs: plūdu karte un rinda sakrīt."
)
eyebrow(s, "Skatpunkts A · Pirms pirkuma")
text(s, 0.8, 0.9, 4.4, 1.5, ["Pārbaudi gabalu", "pirms pērc."], size=32, bold=True)
bullets(
    s,
    0.8,
    2.7,
    4.2,
    [
        "Plūdi 10 / 100 / 200 gadi (LVĢMC)",
        "Daba — ĪADT, mikroliegumi (Ozols)",
        "Apgrūtinājumi (ATIS)",
        "Zonējums (TAPIS)",
    ],
    size=14,
    gap=0.65,
)
video(s, "plusma", 5.3, 0.9, 7.4, 4.625)
text(s, 5.3, 5.7, 7.4, 0.35, "Loki, Spuņciems · 80880050264 · plūdi 100 g. Jā + CFLA 2.1.1.6", size=11, color=MUTED)
footer(s)

# ───────────────────────── 6. Skatpunkts A: slāņi ─────────────────────────
s = new_slide(
    "[4:20–5:00] Tie paši pieci slāņi — Vieta, Kadastrs, Plūdi, Zonējums, Daba. Flīzes nāk tieši no valsts "
    "serveriem. Pircējs redz kontekstu kartē, ne tikai tekstā."
)
eyebrow(s, "Skatpunkts A · Slāņi")
text(s, 0.8, 0.9, 4.3, 1.3, ["Pieci slāņi.", "Viena karte."], size=34, bold=True)
for i, (name, src) in enumerate(
    [
        ("Vieta", "OpenStreetMap / satelīts"),
        ("Kadastrs", "VZD robežas"),
        ("Plūdi", "LVĢMC · 10 / 100 / 200"),
        ("Zonējums", "TAPIS"),
        ("Daba", "Ozols · ĪADT"),
    ]
):
    y = 2.5 + i * 0.72
    text(s, 0.8, y, 1.8, 0.4, name, size=16, bold=True)
    text(s, 2.7, y + 0.02, 2.4, 0.4, src, size=13, color=MUTED)
video(s, "slani", 5.4, 0.85, 7.3, 4.56)
text(s, 5.4, 5.6, 7.3, 0.35, "Karte ↔ Satelīts (apakšā pa kreisi) · Kadastrs · Plūdi · Zonējums · Daba", size=11, color=MUTED)
footer(s)

# ───────────────────────── 7. Skatpunkts B: atbalsts ─────────────────────────
s = new_slide(
    "[5:00–6:20] Skatpunkts B — īpašnieks meklē, kam pieteikties. CFLA atvērtās atlases ar automātisku "
    "atbilstību pēc īpašnieka veida. EKII mājsaimniecību AER — saules, siltumsūkņi. Tikai tas, kas der šim gabalam."
)
eyebrow(s, "Skatpunkts B · Īpašnieks / atbalsts")
text(s, 0.8, 0.9, 4.4, 1.6, ["Ko varu", "pieteikt?"], size=34, bold=True)
bullets(
    s,
    0.8,
    2.75,
    4.2,
    [
        "CFLA 2021–2027 atvērtās atlases",
        "EKII-6 mājsaimniecību AER",
        "Atbilst / Nē pēc īpašnieka veida",
        "Saite uz oficiālo konkursu",
    ],
    size=14,
    gap=0.65,
)
video(s, "atbalsts", 5.3, 0.9, 7.4, 4.625)
text(
    s,
    5.3,
    5.7,
    7.4,
    0.35,
    "Pašvaldība CFLA → Ozoliņi EKII-6 → Vecāķi fiziska + plūdi + EKII",
    size=11,
    color=MUTED,
)
footer(s)

# ───────────────────────── 8. Skatpunkts C: valsts pārvalde ─────────────────────────
s = new_slide(
    "[6:20–7:20] Skatpunkts C — valsts pārvalde. Projektu atlase un finanšu piešķiršana: vienā panelī "
    "īpašnieka veids, jau īstenotie ES projekti uz kadastra, atvērtās atlases, uzņēmums adresē, de minimis. "
    "Nav lēmums — bet vienota fona aina pirms lēmuma."
)
eyebrow(s, "Skatpunkts C · Valsts pārvalde")
text(s, 0.8, 0.9, 4.4, 1.5, ["Viena aina", "pirms atlases."], size=32, bold=True)
bullets(
    s,
    0.8,
    2.6,
    4.2,
    [
        "Īpašnieka veids no VZD",
        "ES projekti uz kadastra",
        "Atvērtās atlases (CFLA)",
        "Uzņēmums: VID, de minimis",
    ],
    size=14,
    gap=0.65,
)
video(s, "uznemums", 5.3, 0.9, 7.4, 4.625)
text(s, 5.3, 5.7, 7.4, 0.35, "Kaļķu 1 + reģ. nr. 40003032949 (Latvenergo) · VID / de minimis", size=11, color=MUTED)
footer(s)

# ───────────────────────── 9. Pierādījums ─────────────────────────
s = new_slide(
    "[7:20–8:00] Katram faktam — Avots un Pieprasījums. Neticiet mums: atveriet oficiālo reģistru vai "
    "precīzo API atbildi. Nekas nav izdomāts."
)
eyebrow(s, "Pierādījums")
text(s, 0.8, 0.9, 4.4, 2.0, ["Neticiet", "mums.", "Pārbaudiet."], size=36, bold=True)
text(
    s,
    0.8,
    3.4,
    4.2,
    1.8,
    ["Avots — oficiālais ieraksts.", "Pieprasījums — tas pats izsaukums, ko redzēja serveris."],
    size=15,
    color=MUTED,
    spacing=8,
)
video(s, "pierad", 5.3, 0.95, 7.4, 4.625)
footer(s)

# ───────────────────────── 10. Meklēšana + telefons ─────────────────────────
s = new_slide(
    "[8:00–8:40] Adrese vai kadastra numurs — karte aizlido. Telefonā — apakšējā loksne. "
    "Desktops, planšete, telefons: viens produkts."
)
eyebrow(s, "Meklēšana un telefons")
text(s, 0.8, 0.85, 8, 0.7, "Raksti. Velc. Klikšķini.", size=32, bold=True)
video(s, "meklesana", 0.8, 1.75, 7.3, 4.56)
text(s, 0.8, 6.45, 7.3, 0.3, "«Kaļķu iela 1, Rīga» → Nominatim → panelis", size=11, color=MUTED)
video(s, "telefons", 9.5, 0.85, 2.95, 6.2)
footer(s)

# ───────────────────────── 11. Arhitektūra ─────────────────────────
s = new_slide(
    "[8:40–9:20] Pārlūks jautā. Lokālais Express tulko. Atvērtie avoti atbild. Katrs avots raksta savu rindu, "
    "tiklīdz atbild — lēnais neaizkavē ātro."
)
eyebrow(s, "Arhitektūra")
title(s, "Pārlūks → serveris → avoti", size=30)


def box(x, y, w, h, head, lines):
    card(s, x, y, w, h)
    text(s, x + 0.25, y + 0.22, w - 0.5, 0.4, head, size=15, bold=True, color=ACCENT)
    text(s, x + 0.25, y + 0.7, w - 0.5, h - 0.9, lines, size=12, color=INK, spacing=4)


box(0.8, 2.15, 3.5, 3.7, "PĀRLŪKS", ["React + Leaflet", "Klikšķis → lat, lon", "Skeleti → rindas", "Slāņu flīzes no valsts"])
box(4.9, 2.15, 3.7, 3.7, "SERVERIS", ["/api/parcel/at", "/api/parcel/report → job", "SQLite kešs 24 h", "cadastre.sqlite · indekss"])
box(9.15, 2.15, 3.4, 3.7, "AVOTI", ["data.gov.lv CKAN", "VZD · TAPIS · Ozols", "CFLA · EKII HTML", "LVĢMC · BIS · UR"])
footer(s)

# ───────────────────────── 12. Kas nav ─────────────────────────
s = new_slide(
    "[9:20–9:45] Ko mēs apzināti nedarām: nav riska punktu, nav personu vārdu, nav ELRG un maksas API. "
    "Automātiskā atbilstība atlasei ir aplēse, ne CFLA vai LVIF lēmums."
)
eyebrow(s, "Robežas")
title(s, "Kas tas nav.", size=36)
for i, (a, b) in enumerate(
    [
        ("Nav lēmums", "Atbilstība ir norāde, ne piešķīrums"),
        ("Nav personas", "Atvērtajā kopā nav vārdu — mēs tos nevācam"),
        ("Nav ELRG", "Tikai atvērtie dati un publiskas HTML tabulas"),
        ("Nav «Nē» no kļūdas", "Ja avots neatbild — to arī rakstām"),
    ]
):
    x = 0.8 + (i % 2) * 6.15
    y = 2.5 + (i // 2) * 2.05
    card(s, x, y, 5.9, 1.85, fill=TINT, line=LINE)
    text(s, x + 0.4, y + 0.35, 5.1, 0.45, a, size=18, bold=True, color=ACCENT_DEEP)
    text(s, x + 0.4, y + 0.95, 5.1, 0.6, b, size=14, color=MUTED)
footer(s)

# ───────────────────────── 13. Noslēgums ─────────────────────────
s = new_slide(
    "[9:45–10:00] Vienotais Zemes Reģistrs. Uzklikšķini zemei. Redzi visu. "
    "Live: hakatons uz 172.82.74.224.sslip.io. Paldies — jautājumi."
)
accent_bar(s, 0.8, 2.4, 0.1, 1.5)
text(s, 1.15, 2.25, 11, 1.0, ["Uzklikšķini zemei.", "Redzi visu."], size=40, bold=True)
text(s, 1.15, 3.6, 10, 0.45, "Paldies.", size=24, color=ACCENT)
card(s, 1.15, 4.5, 10.5, 1.7)
text(
    s,
    1.5,
    4.75,
    9.8,
    1.2,
    [
        "https://172.82.74.224.sslip.io/hakatons/",
        "Dokumentācija: docs/ARHITEKTURA.md",
        "npm run dev  →  localhost:5173",
    ],
    size=15,
    color=MUTED,
    spacing=6,
)
footer(s)

out = HERE / "Vienotais-Zemes-Registrs-portals.pptx"
prs.save(out)
print(f"saved {out}")
